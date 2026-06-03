import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ====================== 【你的配置 - 已直接填好】 ======================
EXCEL_FOLDER = "C:/Users/admin/Downloads/scooter"
OUTPUT_FILE = "全维度分析报告scooter.xlsx"
TARGET_SHEET = "Product Performance"

# 核心列名（完全和你Excel一致）
BRAND_COL = "Brand"
REVENUE_COL = "Revenue"
PRICE_COL = "Avg. Price"
ASIN_COL = "Asin"

TOP_N = 5
# ======================================================================

# 样式定义
TITLE_FONT = Font(name='微软雅黑', size=14, bold=True, color='FFFFFF')
HEADER_FONT = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
DATA_FONT = Font(name='微软雅黑', size=10)
TITLE_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
HEADER_FILL = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
SUM_FILL = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
CENTER_ALIGN = Alignment(horizontal='center', vertical='center')
LEFT_ALIGN = Alignment(horizontal='left', vertical='center')
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

# ---------------------- 1. 品牌分析 ----------------------
def analyze_brand(file_path):
    try:
        df = pd.read_excel(file_path, sheet_name=TARGET_SHEET, engine="openpyxl")
        req = [BRAND_COL, REVENUE_COL, PRICE_COL, ASIN_COL]
        miss = [x for x in req if x not in df.columns]
        if miss:
            return None, f"缺失列：{','.join(miss)}", 0
        
        df = df.dropna(subset=req).copy()
        df[REVENUE_COL] = pd.to_numeric(df[REVENUE_COL], errors="coerce")
        df[PRICE_COL] = pd.to_numeric(df[PRICE_COL], errors="coerce")
        df = df.dropna(subset=[REVENUE_COL, PRICE_COL])
        if df.empty:
            return None, "有效数据为空", 0

        brand_stats = df.groupby(BRAND_COL, as_index=False).agg(
            总REVENUE=(REVENUE_COL, "sum"),
            平均售价=(PRICE_COL, "mean")
        )
        total_rev = brand_stats["总REVENUE"].sum()
        brand_stats["REVENUE占比(%)"] = (brand_stats["总REVENUE"] / total_rev * 100).round(2)

        asin_total = df.groupby([BRAND_COL, ASIN_COL], as_index=False)[REVENUE_COL].sum()
        asin_total = asin_total.sort_values([BRAND_COL, REVENUE_COL], ascending=[True, False])
        top_asin_per_brand = asin_total.groupby(BRAND_COL).head(1)
        top_asin_per_brand.columns = [BRAND_COL, ASIN_COL, "该ASIN_REVENUE"]

        res = pd.merge(brand_stats, top_asin_per_brand, on=BRAND_COL)
        res = res.sort_values("总REVENUE", ascending=False).head(TOP_N).reset_index(drop=True)
        top5_sum = res["REVENUE占比(%)"].sum().round(2)
        return res, "成功", top5_sum

    except Exception as e:
        return None, f"分析失败：{str(e)}", 0

# ---------------------- 2. ASIN分析（已修复：按ASIN汇总Revenue） ----------------------
def analyze_asin(file_path):
    try:
        df = pd.read_excel(file_path, sheet_name=TARGET_SHEET, engine="openpyxl")
        req = [BRAND_COL, REVENUE_COL, PRICE_COL, ASIN_COL]
        miss = [x for x in req if x not in df.columns]
        if miss:
            return None, f"缺失列：{','.join(miss)}", 0
        
        df = df.dropna(subset=req).copy()
        df[REVENUE_COL] = pd.to_numeric(df[REVENUE_COL], errors="coerce")
        df[PRICE_COL] = pd.to_numeric(df[PRICE_COL], errors="coerce")
        df = df.dropna(subset=[REVENUE_COL])
        if df.empty:
            return None, "有效数据为空", 0

        # ✅ 核心：相同Asin的Revenue自动求和
        asin_summary = df.groupby(ASIN_COL, as_index=False).agg(
            总REVENUE=(REVENUE_COL, "sum"),
            平均售价=(PRICE_COL, "mean"),
            Brand=(BRAND_COL, "first")
        )

        total_rev = asin_summary["总REVENUE"].sum()
        asin_summary["REVENUE占比(%)"] = (asin_summary["总REVENUE"] / total_rev * 100).round(2)
        top5 = asin_summary.sort_values("总REVENUE", ascending=False).head(TOP_N).reset_index(drop=True)
        top5_sum = top5["REVENUE占比(%)"].sum().round(2)

        return top5, "成功", top5_sum

    except Exception as e:
        return None, f"分析失败：{str(e)}", 0

# ---------------------- 写入品牌表 ----------------------
def write_brand_sheet(ws, files):
    current_row = 1
    headers = ["Brand", "总REVENUE", "平均售价", "REVENUE占比(%)", "最高REVENUE_ASIN", "该ASIN_REVENUE"]
    col_cnt = len(headers)
    
    for file in files:
        res, msg, total_ratio = analyze_brand(file)
        title = f"文件：{file.name} | {msg} | TOP5品牌总占比：{total_ratio}%" if msg == "成功" else f"文件：{file.name} | {msg}"
        
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_cnt)
        c = ws.cell(row=current_row, column=1, value=title)
        c.font = TITLE_FONT; c.fill = TITLE_FILL; c.alignment = CENTER_ALIGN; c.border = THIN_BORDER
        current_row += 1

        for i, h in enumerate(headers, 1):
            c = ws.cell(row=current_row, column=i, value=h)
            c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = CENTER_ALIGN; c.border = THIN_BORDER
        current_row += 1

        if res is not None:
            for _, r in res.iterrows():
                ws.cell(row=current_row, column=1, value=r[BRAND_COL]).alignment = LEFT_ALIGN
                ws.cell(row=current_row, column=2, value=r["总REVENUE"]).alignment = CENTER_ALIGN
                ws.cell(row=current_row, column=3, value=round(r["平均售价"], 2)).alignment = CENTER_ALIGN
                ws.cell(row=current_row, column=4, value=r["REVENUE占比(%)"]).alignment = CENTER_ALIGN
                ws.cell(row=current_row, column=5, value=r[ASIN_COL]).alignment = LEFT_ALIGN
                ws.cell(row=current_row, column=6, value=r["该ASIN_REVENUE"]).alignment = CENTER_ALIGN
                for col in range(1, col_cnt+1):
                    cell = ws.cell(row=current_row, column=col)
                    cell.font = DATA_FONT; cell.border = THIN_BORDER
                current_row += 1

            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_cnt)
            s = ws.cell(row=current_row, column=1, value=f"🔥 TOP5品牌总营收占比：{total_ratio}%")
            s.font = Font(bold=True); s.fill = SUM_FILL; s.alignment = CENTER_ALIGN; s.border = THIN_BORDER
            current_row += 1
        else:
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_cnt)
            ws.cell(row=current_row, column=1, value=msg).font = DATA_FONT
            current_row += 1
        current_row += 1

    for i in range(1, col_cnt+1):
        max_len = 0
        for row in range(1, ws.max_row + 1):
            val = str(ws.cell(row=row, column=i).value or "")
            max_len = max(max_len, sum(2 if ord(c) > 127 else 1 for c in val))
        ws.column_dimensions[get_column_letter(i)].width = max_len + 2

# ---------------------- 写入ASIN表 ----------------------
def write_asin_sheet(ws, files):
    current_row = 1
    headers = ["Asin", "Brand", "总REVENUE", "平均售价", "REVENUE占比(%)"]
    col_cnt = len(headers)
    
    for file in files:
        res, msg, total_ratio = analyze_asin(file)
        title = f"文件：{file.name} | {msg} | TOP5 ASIN总占比：{total_ratio}%" if msg == "成功" else f"文件：{file.name} | {msg}"
        
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_cnt)
        c = ws.cell(row=current_row, column=1, value=title)
        c.font = TITLE_FONT; c.fill = TITLE_FILL; c.alignment = CENTER_ALIGN; c.border = THIN_BORDER
        current_row += 1

        for i, h in enumerate(headers, 1):
            c = ws.cell(row=current_row, column=i, value=h)
            c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = CENTER_ALIGN; c.border = THIN_BORDER
        current_row += 1

        if res is not None:
            for _, r in res.iterrows():
                ws.cell(row=current_row, column=1, value=r[ASIN_COL]).alignment = LEFT_ALIGN
                ws.cell(row=current_row, column=2, value=r["Brand"]).alignment = LEFT_ALIGN
                ws.cell(row=current_row, column=3, value=r["总REVENUE"]).alignment = CENTER_ALIGN
                ws.cell(row=current_row, column=4, value=round(r["平均售价"], 2)).alignment = CENTER_ALIGN
                ws.cell(row=current_row, column=5, value=r["REVENUE占比(%)"]).alignment = CENTER_ALIGN
                for col in range(1, col_cnt+1):
                    cell = ws.cell(row=current_row, column=col)
                    cell.font = DATA_FONT; cell.border = THIN_BORDER
                current_row += 1

            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_cnt)
            s = ws.cell(row=current_row, column=1, value=f"🔥 TOP5 ASIN总营收占比：{total_ratio}%")
            s.font = Font(bold=True); s.fill = SUM_FILL; s.alignment = CENTER_ALIGN; s.border = THIN_BORDER
            current_row += 1
        else:
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_cnt)
            ws.cell(row=current_row, column=1, value=msg).font = DATA_FONT
            current_row += 1
        current_row += 1

    for i in range(1, col_cnt+1):
        max_len = 0
        for row in range(1, ws.max_row + 1):
            val = str(ws.cell(row=row, column=i).value or "")
            max_len = max(max_len, sum(2 if ord(c) > 127 else 1 for c in val))
        ws.column_dimensions[get_column_letter(i)].width = max_len + 2

# ---------------------- 主程序 ----------------------
if __name__ == "__main__":
    folder = Path(EXCEL_FOLDER)
    if not folder.exists():
        print("❌ 文件夹不存在")
        exit()
    files = sorted(list(folder.glob("*.xlsx")) + list(folder.glob("*.xls")))
    if not files:
        print("❌ 未找到Excel文件")
        exit()

    from openpyxl import Workbook
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "品牌分析TOP5"
    write_brand_sheet(ws1, files)

    ws2 = wb.create_sheet("ASIN营收TOP5")
    write_asin_sheet(ws2, files)

    wb.save(OUTPUT_FILE)
    print(f"\n✅ 分析完成！结果已保存至：{OUTPUT_FILE}")
